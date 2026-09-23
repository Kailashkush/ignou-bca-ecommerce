"""
Word document construction helpers for the BCA project report.

IGNOU requires the report on A4, double spaced, hard bound, with a contents
page and every page numbered. All of that is configured once here so the
content modules only describe *what* to say, never how to format it.
"""
import pathlib
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = pathlib.Path(__file__).parent

BODY_FONT = 'Times New Roman'
MONO_FONT = 'Consolas'
INK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x44, 0x44, 0x44)
ACCENT = RGBColor(0x1F, 0x3B, 0x73)


# --------------------------------------------------------------------------
# Low-level XML helpers (python-docx exposes no API for these)
# --------------------------------------------------------------------------
def _field(paragraph, instruction):
    """Inserts a Word field code, e.g. PAGE or TOC."""
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t'); placeholder.text = '…'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for node in (begin, instr, sep, placeholder, end):
        run._r.append(node)
    return run


def _shade(cell, hex_colour):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_colour)
    cell._tc.get_or_add_tcPr().append(shd)


def _cell_margins(table, top=60, bottom=60, left=110, right=110):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for tag, value in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        node = OxmlElement(f'w:{tag}')
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')
        mar.append(node)
    tblPr.append(mar)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    node = OxmlElement('w:tblHeader')
    node.set(qn('w:val'), 'true')
    trPr.append(node)


def _keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


# --------------------------------------------------------------------------
# Report builder
# --------------------------------------------------------------------------
class Report:
    def __init__(self):
        self.doc = Document()
        self.figure_no = 0
        self.table_no = 0
        self.figures = []   # (label, caption) — report body only
        self.tables = []
        # Attachments (the synopsis, the appendix) carry their own numbering
        # series and are deliberately absent from the report's index pages.
        self.aux = False
        self.aux_figure_no = 0
        self.aux_table_no = 0
        self.ref_index = {}   # populated from a first pass; see ref()
        self._setup_styles()
        self._setup_page()

    # -- setup -------------------------------------------------------------
    def _setup_page(self):
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)      # A4
        section.page_height = Cm(29.7)
        section.left_margin = Inches(1.4)  # extra allowance for hard binding
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    def _setup_styles(self):
        styles = self.doc.styles

        normal = styles['Normal']
        normal.font.name = BODY_FONT
        normal.font.size = Pt(12)
        normal.font.color.rgb = INK
        normal._element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE   # IGNOU requirement
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)

        for name, size, before, after in (
                ('Heading 1', 16, 12, 8),
                ('Heading 2', 14, 11, 6),
                ('Heading 3', 12.5, 9, 5)):
            style = styles[name]
            style.font.name = BODY_FONT
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = ACCENT
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            style.paragraph_format.keep_with_next = True

    def add_page_numbers(self):
        """Footer showing 'Page N' on every page of the body."""
        for section in self.doc.sections:
            footer = section.footer
            p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = p.add_run('Page ')
            run.font.size = Pt(9); run.font.name = BODY_FONT; run.font.color.rgb = GREY
            _field(p, ' PAGE ')
            for r in p.runs:
                r.font.size = Pt(9); r.font.name = BODY_FONT; r.font.color.rgb = GREY

    def add_header_text(self, text):
        for section in self.doc.sections:
            header = section.header
            p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = p.add_run(text)
            run.font.size = Pt(9); run.font.name = BODY_FONT
            run.font.color.rgb = GREY; run.italic = True

    # -- block elements -----------------------------------------------------
    def h1(self, text, page_break=True):
        if page_break:
            self.page_break()
        p = self.doc.add_heading(text, level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return p

    def h2(self, text):
        return self.doc.add_heading(text, level=2)

    def h3(self, text):
        return self.doc.add_heading(text, level=3)

    def p(self, text, justify=True, first_line_indent=True, spacing='double',
          italic=False, bold=False, size=12, align=None, space_after=0):
        para = self.doc.add_paragraph()
        pf = para.paragraph_format
        if spacing == 'double':
            pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        elif spacing == 'single':
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        else:
            pf.line_spacing = spacing
        pf.space_after = Pt(space_after)
        if align is not None:
            para.alignment = align
        elif justify:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if first_line_indent and justify and align is None:
            pf.first_line_indent = Inches(0.4)
        run = para.add_run(text)
        run.font.size = Pt(size)
        run.italic = italic
        run.bold = bold
        return para

    def bullets(self, items, style='List Bullet', spacing='single'):
        for item in items:
            para = self.doc.add_paragraph(style=style)
            pf = para.paragraph_format
            pf.line_spacing_rule = (WD_LINE_SPACING.SINGLE if spacing == 'single'
                                    else WD_LINE_SPACING.DOUBLE)
            pf.space_after = Pt(6)
            pf.left_indent = Inches(0.55)
            self._rich(para, item)
        self.spacer(4)

    def numbered(self, items, spacing='single'):
        self.bullets(items, style='List Number', spacing=spacing)

    def _rich(self, para, text, size=11.5):
        """Renders **bold** and `code` markers inside a paragraph."""
        import re
        for chunk in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
            if not chunk:
                continue
            if chunk.startswith('**') and chunk.endswith('**'):
                run = para.add_run(chunk[2:-2]); run.bold = True; run.font.size = Pt(size)
            elif chunk.startswith('`') and chunk.endswith('`'):
                run = para.add_run(chunk[1:-1])
                run.font.name = MONO_FONT; run.font.size = Pt(size - 1.5)
            else:
                run = para.add_run(chunk); run.font.size = Pt(size)

    def spacer(self, points=8):
        para = self.doc.add_paragraph()
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        para.paragraph_format.space_after = Pt(points)
        for run in para.runs:
            run.font.size = Pt(2)
        return para

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # -- figures and tables -------------------------------------------------
    def figure(self, filename, caption, width_inches=6.1, folder='diagrams'):
        if self.aux:
            self.aux_figure_no += 1
        else:
            self.figure_no += 1
        path = BASE / folder / filename
        if not path.exists():
            raise FileNotFoundError(path)

        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        para.paragraph_format.space_before = Pt(10)
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(str(path), width=Inches(width_inches))

        label = (f'Figure S{self.aux_figure_no}' if self.aux
                 else f'Figure {self.figure_no}')
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        cap.paragraph_format.space_after = Pt(14)
        run = cap.add_run(f'{label}: {caption}')
        run.font.size = Pt(10); run.italic = True; run.font.color.rgb = GREY
        if not self.aux:
            self.figures.append((label, caption))
        return label

    def screenshot(self, filename, caption, width_inches=6.1):
        return self.figure(filename, caption, width_inches, folder='screenshots')


    def figure_row(self, entries, folder='screenshots', gap_caption=None):
        """Places several images side by side in one borderless row, with a
        numbered caption beneath each. Used for the responsive screenshots,
        where three tall phone captures would otherwise consume a page each."""
        from docx.enum.table import WD_TABLE_ALIGNMENT
        table = self.doc.add_table(rows=2, cols=len(entries))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        labels = []
        for i, (filename, caption, width) in enumerate(entries):
            if self.aux:
                self.aux_figure_no += 1
                label = f'Figure S{self.aux_figure_no}'
            else:
                self.figure_no += 1
                label = f'Figure {self.figure_no}'
            labels.append(label)
            path = BASE / folder / filename
            cell = table.rows[0].cells[i]
            cell.width = Inches(6.1 / len(entries))
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            para.paragraph_format.space_after = Pt(2)
            para.add_run().add_picture(str(path), width=Inches(width))

            cap_cell = table.rows[1].cells[i]
            cap_cell.width = Inches(6.1 / len(entries))
            cpara = cap_cell.paragraphs[0]
            cpara.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cpara.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            cpara.paragraph_format.space_after = Pt(10)
            run = cpara.add_run(f'{label}: {caption}')
            run.font.size = Pt(9); run.italic = True; run.font.color.rgb = GREY
            if not self.aux:
                self.figures.append((label, caption))
        return labels

    def table(self, headers, rows, caption=None, widths=None, font_size=9.5,
              header_fill='DCE3F0', zebra=True):
        if self.aux:
            self.aux_table_no += 1
        else:
            self.table_no += 1
        if caption:
            cap = self.doc.add_paragraph()
            cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            cap.paragraph_format.space_before = Pt(10)
            cap.paragraph_format.space_after = Pt(4)
            cap.paragraph_format.keep_with_next = True
            label = (f'Table S{self.aux_table_no}' if self.aux
                     else f'Table {self.table_no}')
            run = cap.add_run(f'{label}: {caption}')
            run.font.size = Pt(10); run.bold = True; run.font.color.rgb = ACCENT
            if not self.aux:
                self.tables.append((label, caption))

        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _cell_margins(table)

        hdr = table.rows[0]
        _repeat_header(hdr)
        for i, text in enumerate(headers):
            cell = hdr.cells[i]
            cell.text = ''
            para = cell.paragraphs[0]
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(str(text))
            run.bold = True; run.font.size = Pt(font_size); run.font.name = BODY_FONT
            _shade(cell, header_fill)

        for r, row in enumerate(rows):
            cells = table.add_row().cells
            for i, text in enumerate(row):
                cell = cells[i]
                cell.text = ''
                para = cell.paragraphs[0]
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                para.paragraph_format.space_after = Pt(0)
                self._rich(para, str(text), size=font_size)
                for run in para.runs:
                    if run.font.name != MONO_FONT:
                        run.font.name = BODY_FONT
                if zebra and r % 2 == 1:
                    _shade(cell, 'F4F6FA')

        if widths:
            total = sum(widths)
            usable = Inches(6.1)
            for row in table.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(6.1 * w / total)

        self.spacer(12)
        return table

    def code(self, text, caption=None, size=8.0, language=None):
        """Monospaced, single-spaced listing inside a bordered single-cell table."""
        if caption:
            cap = self.doc.add_paragraph()
            cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            cap.paragraph_format.space_before = Pt(10)
            cap.paragraph_format.space_after = Pt(3)
            cap.paragraph_format.keep_with_next = True
            run = cap.add_run(caption)
            run.font.size = Pt(9.5); run.bold = True; run.font.color.rgb = ACCENT

        table = self.doc.add_table(rows=1, cols=1)
        table.style = 'Table Grid'
        table.autofit = False
        _cell_margins(table, top=80, bottom=80, left=140, right=100)
        cell = table.rows[0].cells[0]
        cell.width = Inches(6.1)
        cell.text = ''
        _shade(cell, 'FBFBFD')

        lines = text.rstrip('\n').split('\n')
        for i, line in enumerate(lines):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            pf = para.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            run = para.add_run(line if line else ' ')
            run.font.name = MONO_FONT
            run.font.size = Pt(size)
            run._element.rPr.rFonts.set(qn('w:eastAsia'), MONO_FONT)
        self.spacer(10)
        return table

    def toc_field(self, levels='1-3'):
        para = self.doc.add_paragraph()
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        _field(para, f' TOC \\o "{levels}" \\h \\z \\u ')

    def centered(self, text, size=12, bold=False, italic=False, space_after=6,
                 caps=False, colour=None):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        para.paragraph_format.space_after = Pt(space_after)
        run = para.add_run(text.upper() if caps else text)
        run.font.size = Pt(size); run.bold = bold; run.italic = italic
        if colour is not None:
            run.font.color.rgb = colour
        return para


    # -- symbolic cross-references -----------------------------------------
    # Prose refers to a figure or table by a distinctive fragment of its
    # caption rather than by a number. The number is resolved from the
    # registry built by a first pass, so inserting a figure in Chapter 3 can
    # never silently leave a stale "see Table 13" in Chapter 5. An unresolved
    # reference raises at build time instead of shipping.
    def ref(self, kind, caption_fragment):
        key = caption_fragment.lower()
        matches = [label for (k, cap), label in self.ref_index.items()
                   if k == kind and key in cap]
        if len(matches) == 1:
            return matches[0]
        if not self.ref_index:
            return f'{kind.capitalize()} ?'          # first pass
        if not matches:
            raise KeyError(f'cross-reference matches no {kind}: {caption_fragment!r}')
        raise KeyError(
            f'cross-reference is ambiguous for {kind} {caption_fragment!r} -> {matches}')

    def ref_table(self, caption_fragment):
        return self.ref('table', caption_fragment)

    def ref_figure(self, caption_fragment):
        return self.ref('figure', caption_fragment)

    def make_ref_index(self):
        """Caption -> label map, for feeding into the second pass."""
        index = {}
        for label, caption in self.figures:
            index[('figure', caption.lower())] = label
        for label, caption in self.tables:
            index[('table', caption.lower())] = label
        return index

    def right_text(self, text, size=11.5, bold=True):
        """Right-aligned block, used for the signature line beneath the
        acknowledgement. Newlines become line breaks."""
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = para.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        return para

    def rule(self):
        para = self.doc.add_paragraph()
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pPr = para._p.get_or_add_pPr()
        borders = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '8')
        bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '1F3B73')
        borders.append(bottom)
        pPr.append(borders)
        return para

    def save(self, path):
        self.doc.save(str(path))
        return path
